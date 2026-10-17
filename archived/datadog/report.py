"""Build a fillable PFI working draft from DataDog's structured evidence."""

from __future__ import annotations

import argparse
from datetime import datetime, timezone
import io
import json
from pathlib import Path
import sys

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image


NAVY = "233B5D"
PALE_BLUE = "EAF0F6"
PALE_GRAY = "F3F5F7"
GRID = "D5DCE4"
BLACK = "000000"
GRAY = "58616D"
MAX_EMBEDDED_PHOTOS = 80


def shade(cell, fill: str) -> None:
    props = cell._tc.get_or_add_tcPr()
    shd = props.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        props.append(shd)
    shd.set(qn("w:fill"), fill)


def borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    edges = tbl_pr.first_child_found_in("w:tblBorders")
    if edges is None:
        edges = OxmlElement("w:tblBorders")
        tbl_pr.append(edges)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = edges.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            edges.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "5")
        element.set(qn("w:color"), GRID)


def set_cell_margins(cell, top=100, start=110, bottom=100, end=110) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn("w:" + margin))
        if node is None:
            node = OxmlElement("w:" + margin)
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_text(cell, value: str, *, bold=False, color=BLACK, size=9) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.08
    run = p.add_run(str(value or ""))
    run.bold = bold
    run.font.name = "Arial"
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)


def add_table(doc, headers: list[str], rows: list[list[str]], widths: list[float] | None = None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    repeat_header(table.rows[0])
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        shade(cell, NAVY)
        set_cell_text(cell, header, bold=True, color="FFFFFF", size=8.5)
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for idx, value in enumerate(values):
            if row_index % 2 == 1:
                shade(cells[idx], PALE_GRAY)
            set_cell_text(cells[idx], str(value), size=8.5)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            if widths and idx < len(widths):
                cell.width = Inches(widths[idx])
    borders(table)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_section(doc, title: str, note: str | None = None) -> None:
    p = doc.add_paragraph(style="Heading 1")
    p.paragraph_format.keep_with_next = True
    p.add_run(title)
    if note:
        p = doc.add_paragraph(note)
        p.paragraph_format.space_after = Pt(7)


def set_page_field(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    run.font.name = "Arial"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(GRAY)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    value = OxmlElement("w:t")
    value.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for item in (begin, instruction, separate, value, end):
        run._r.append(item)


def format_document(doc) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(9.5)
    normal.font.color.rgb = RGBColor.from_string(BLACK)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.12
    for name, size in (("Title", 23), ("Heading 1", 14), ("Heading 2", 11)):
        style = styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(BLACK)
        style.paragraph_format.space_before = Pt(12 if name != "Title" else 0)
        style.paragraph_format.space_after = Pt(6)
        p_pr = style._element.get_or_add_pPr()
        p_borders = p_pr.find(qn("w:pBdr"))
        if p_borders is not None:
            p_pr.remove(p_borders)
    footer = section.footer.paragraphs[0]
    footer.text = "DataDog working draft  |  Complete missing particulars and obtain CP review"
    footer.paragraph_format.space_before = Pt(3)
    footer.runs[0].font.name = "Arial"
    footer.runs[0].font.size = Pt(7.5)
    footer.runs[0].font.color.rgb = RGBColor.from_string(GRAY)
    set_page_field(section.footer.add_paragraph())


def blank(value: object) -> str:
    if value is None or str(value).strip() == "":
        return "[To be completed]"
    return str(value)


def save_photo_bytes(path: Path) -> io.BytesIO:
    image = Image.open(path).convert("RGB")
    image.thumbnail((1400, 1400))
    output = io.BytesIO()
    image.save(output, format="JPEG", quality=82, optimize=True)
    output.seek(0)
    return output


def create_report(data: dict, output: Path) -> None:
    doc = Document()
    format_document(doc)
    project = data.get("project", {})
    evidence = data.get("evidence", [])
    elevations = data.get("elevations", [])
    descriptions = data.get("descriptions", [])
    location = project.get("location") or "[Building name]"
    generated = datetime.fromtimestamp(data.get("generated_at", 0), timezone.utc).astimezone()

    title = doc.add_paragraph(style="Title")
    title.add_run("Periodic Façade Inspection Draft Report")
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(10)
    r = subtitle.add_run(f"{location}  |  Generated {generated.strftime('%d %b %Y %H:%M %Z')}")
    r.font.name = "Arial"
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor.from_string(GRAY)
    doc.add_paragraph(
        "Working draft compiled from DataDog records. Bracketed fields are incomplete. "
        "AI observations are prompts for human review, not confirmed defects or the "
        "Competent Person’s assessment. Complete and verify this document against the "
        "current BCA report template before submission."
    )

    add_section(doc, "1 Building particulars")
    particulars = [
        ["Building or development", blank(location), "Address", "[To be completed]"],
        ["Block number", "[To be completed]", "Building use", "[To be completed]"],
        ["TOP or CSC date", "[To be completed]", "Building age", "[To be completed]"],
        ["Building height", "[To be completed]", "Number of storeys", "[To be completed]"],
        ["Owner or MCST", "[To be completed]", "BCA notice number", "[To be completed]"],
        ["Report reference", "[To be completed]", "Previous inspection date", "[To be completed]"],
    ]
    add_table(doc, ["Particular", "Details", "Particular", "Details"], particulars, [1.35, 2.15, 1.35, 2.0])

    add_section(doc, "2 Inspection personnel and method")
    add_table(doc, ["Item", "Details to complete"], [
        ["Competent Person and registration", "[Name, company, registration and contact details]"],
        ["Façade Inspector and registration", "[Name, company and registration number]"],
        ["Inspection dates and weather", "[Dates, times and weather conditions]"],
        ["Visual inspection method", "[Camera, vantage points, drone/UAS or other approved method]"],
        ["Close-range access and equipment", "[Rope access, gondola, boom lift, tools and equipment]"],
        ["Risk and work-at-height records", "[Risk assessment, method statement, PTW and rescue plan references]"],
    ], [2.3, 4.55])

    add_section(doc, "3 Drawings and façade information")
    add_table(doc, ["Record", "Status or reference"], [
        ["Approved plans, floor layouts and elevation drawings", "[To be completed]"],
        ["Façade systems, materials and connections", "[To be completed]"],
        ["Previous PFI reports and defect records", "[To be completed]"],
        ["Maintenance and repair history", "[To be completed]"],
        ["Known complaints or reported façade issues", "[To be completed]"],
    ], [3.55, 3.3])

    add_section(doc, "4 Inspection checklist and evidence status", "A recorded photo shows that evidence was submitted. It does not by itself establish inspection coverage or completion.")
    checklist_rows = [
        ["Review approved drawings, prior reports and repair history", "[Record reviewed documents]"],
        ["Visual inspection of the full façade", "[Confirm coverage and any inaccessible areas]"],
        ["Overall building overview", "Recorded" if any(e.get("label") == "building_overview" for e in evidence) else "Missing"],
        ["Plan drawing", "Recorded" if any(e.get("label") == "plan" for e in evidence) else "Missing"],
        ["Overall photo of each elevation", "See elevation rows below"],
        ["Medium-long photos, not more than four storeys per image", "[Confirm images and coverage]"],
        ["Close-range inspection of at least 10% of each elevation", "[Enter area calculation for each elevation]"],
        ["Close-range locations marked on elevation drawings", "[Attach marked drawings]"],
        ["Inspection methods, dates, tools and personnel recorded", "[Complete section 2 and the elevation records]"],
        ["Defect photos, location references and field records", f"{len(evidence)} evidence record(s) in this draft"],
    ]
    for elevation in elevations:
        name = elevation.get("name", "Unnamed elevation")
        photos = [e for e in evidence if e.get("elevation") == name and e.get("label") == "elevation_overview"]
        checklist_rows.append([f"Elevation {name}: overall photo", "Recorded" if photos else "Missing"])
        checklist_rows.append([f"Elevation {name}: 10% close-range area and calculation", "[Enter total elevation area and inspected area]"])
        for finding in elevation.get("findings", []):
            number = finding.get("number")
            finding_id = f"F{number:03}" if isinstance(number, int) else "[Finding]"
            finding_evidence = [e for e in evidence if e.get("finding") == finding_id]
            for label, title in (("context_photo", "context photo"), ("close_up", "close-up photo")):
                present = any(e.get("label") == label for e in finding_evidence)
                checklist_rows.append([f"{finding_id}: {title}", "Recorded" if present else "Missing"])
            has_description = any(e.get("description") for e in finding_evidence) or any(
                d.get("finding") == finding_id and d.get("text") for d in descriptions
            )
            checklist_rows.append([f"{finding_id}: description", "Recorded" if has_description else "Missing"])
    add_table(doc, ["PFI item", "Evidence or action required"], checklist_rows, [3.65, 3.2])

    add_section(doc, "5 Elevation records")
    if elevations:
        rows = []
        for elevation in elevations:
            rows.append([
                elevation.get("name", "[Name]"),
                "[Enter orientation, dimensions and façade system]",
                "[Enter close-range area and calculation]",
                "[Enter dates, duration, method and tools]",
            ])
        add_table(doc, ["Elevation", "Description", "Close-range coverage", "Inspection record"], rows, [1.0, 2.05, 1.8, 2.0])
    else:
        doc.add_paragraph("[Add each building face and its reference, dimensions, façade system and inspection record.]")

    add_section(doc, "6 Preliminary visual observations")
    observations = []
    for item in evidence:
        for observation in item.get("observations") or []:
            observations.append([
                f"E{item.get('id', 0):04}",
                blank(item.get("elevation")) + (f" / {item['zone']}" if item.get("zone") else ""),
                f"{observation.get('component', 'Unknown component')} — {observation.get('condition', 'condition to review')}",
                observation.get("note", "[No observation note]"),
                ", ".join(observation.get("verification") or []) or "CP/FI review",
            ])
    if observations:
        add_table(doc, ["Evidence", "Location", "AI candidate", "Visible observation", "Review or verification"], observations, [0.7, 1.15, 1.3, 2.45, 1.25])
    else:
        doc.add_paragraph("No AI visual observations are recorded. Add observations after reviewing the submitted photographs.")
    doc.add_paragraph("All entries in this section are automated visual candidates. Confirm the component, condition and location on site. Physical tests and the CP’s engineering judgement must be recorded separately.")

    add_section(doc, "7 Field checks and inspector notes")
    field_rows = []
    for item in evidence:
        if item.get("field_note"):
            field_rows.append([
                f"E{item.get('id', 0):04}",
                blank(item.get("elevation")) + (f" / {item['zone']}" if item.get("zone") else ""),
                item["field_note"],
            ])
    if field_rows:
        add_table(doc, ["Evidence", "Location", "Recorded field note"], field_rows, [0.8, 1.65, 4.4])
    else:
        doc.add_paragraph("[Record tapping, probing, borescope, adhesion, carbonation or other inspection results with method, date, location and inspector.]")

    add_section(doc, "8 Defect schedule")
    defect_rows = []
    for item in evidence:
        obs = item.get("observations") or []
        if item.get("description") or item.get("field_note") or obs:
            summary = "; ".join(o.get("note", "") for o in obs if o.get("note"))
            defect_rows.append([
                item.get("finding") or f"E{item.get('id', 0):04}",
                blank(item.get("elevation")) + (f" / {item['zone']}" if item.get("zone") else ""),
                "; ".join(filter(None, [item.get("description"), summary, item.get("field_note")])) or "[Complete observation]",
                "[CP to classify and recommend action]",
            ])
    for description in descriptions:
        finding_id = description.get("finding")
        matching = next((row for row in defect_rows if row[0] == finding_id), None)
        if matching:
            if description.get("text") and description["text"] not in matching[2]:
                matching[2] += "; " + description["text"]
        elif description.get("text"):
            defect_rows.append([
                finding_id or "[Finding]",
                "[Location to complete]",
                description["text"],
                "[CP to classify and recommend action]",
            ])
    if defect_rows:
        add_table(doc, ["Finding", "Location", "Recorded observation", "CP assessment and action"], defect_rows, [0.75, 1.3, 3.25, 1.55])
    else:
        add_table(doc, ["Finding", "Location", "Recorded observation", "CP assessment and action"], [["[ID]", "[Elevation, grid and level]", "[Condition, type and extent]", "[Complete]" ]], [0.75, 1.3, 3.25, 1.55])

    add_section(doc, "9 Competent Person assessment")
    doc.add_paragraph("The appointed CP records the condition classification for each façade system after reviewing the complete inspection evidence.")
    add_table(doc, ["Façade system", "Condition: Safe / Require Repair / Unsafe", "Assessment and basis"], [["[System and location]", "[CP to complete]", "[CP to complete]"], ["[System and location]", "[CP to complete]", "[CP to complete]" ]], [1.8, 2.2, 2.85])

    add_section(doc, "10 Recommendations and follow-up")
    add_table(doc, ["Finding or façade area", "Recommended action by CP", "Precaution, owner and target date"], [["[Complete]", "[Complete]", "[Complete]"], ["[Complete]", "[Complete]", "[Complete]" ]], [1.6, 2.7, 2.55])

    add_section(doc, "11 Conclusion and further investigation")
    doc.add_paragraph("Overall assessment: [To be completed by the appointed CP]")
    doc.add_paragraph("Further façade investigation required: [Yes / No, with basis and approval reference where applicable]")
    doc.add_paragraph("Immediate precautionary measures: [To be completed by the CP]")

    add_section(doc, "12 Photo register")
    photo_rows = []
    for item in evidence:
        photo_rows.append([
            f"E{item.get('id', 0):04}",
            datetime.fromtimestamp(item.get("sent_at", 0), timezone.utc).astimezone().strftime("%d %b %Y %H:%M"),
        (item.get("label") or "Unlabelled").replace("_", " ").title(),
            "/".join(filter(None, [item.get("elevation"), item.get("zone"), item.get("finding")] )) or "[Location needed]",
            item.get("filename") or "[Telegram image]",
        ])
    if photo_rows:
        add_table(doc, ["Evidence", "Received", "Photo role", "Location", "Original file"], photo_rows, [0.7, 1.0, 1.05, 1.7, 2.4])
    else:
        doc.add_paragraph("No project evidence photos are currently recorded.")

    add_section(doc, "13 Competent Person certification")
    doc.add_paragraph("Insert the current required certification wording. The appointed CP must review the completed report, then sign, stamp and date it.")
    add_table(doc, ["Name and registration", "Signature and stamp", "Date"], [["[Complete]", "", ""]], [2.25, 2.8, 1.8])

    add_section(doc, "Photo annex")
    doc.add_paragraph("Original evidence remains stored in the project archive. The images below are report copies of up to 80 available photos; refer to the photo register for all evidence IDs.")
    embedded = 0
    for item in evidence:
        image_path = item.get("preview_path") or item.get("source_path")
        if not image_path or embedded >= MAX_EMBEDDED_PHOTOS:
            continue
        path = Path(image_path)
        if not path.is_file():
            continue
        caption = f"E{item.get('id', 0):04} | {item.get('label') or 'Unlabelled'} | {item.get('elevation') or 'Location not recorded'}"
        if item.get("zone"):
            caption += f" | {item['zone']}"
        p = doc.add_paragraph()
        p.paragraph_format.keep_with_next = True
        run = p.add_run(caption)
        run.bold = True
        run.font.name = "Arial"
        run.font.size = Pt(9)
        try:
            doc.add_picture(save_photo_bytes(path), width=Inches(5.8))
        except Exception:
            doc.add_paragraph(f"[Image unavailable: {path.name}]")
        embedded += 1
    if len([e for e in evidence if e.get("source_path")]) > MAX_EMBEDDED_PHOTOS:
        doc.add_paragraph("Some images are listed in the photo register but not embedded because the report annex is limited to 80 photos. The originals remain in the evidence archive.")

    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", required=True, type=Path)
    args = parser.parse_args()
    data = json.load(sys.stdin)
    create_report(data, args.output)


if __name__ == "__main__":
    main()
